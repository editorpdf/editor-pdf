"""Ferramentas em lote do Editor de PDF (estilo iLovePDF).

Cada ferramenta recebe uma lista de arquivos [(nome, bytes)] e um dicionário de opções,
e devolve [(nome, bytes)] mais uma lista de avisos. O servidor chama executar().
Tudo roda no computador onde o servidor está ligado; nada vai para a internet
(exceto "HTML para PDF" quando você pede uma página da web).
"""
import difflib
import io
import ipaddress
import os
import re
import shutil
import socket
import subprocess
import tempfile
import textwrap
import threading
import uuid
import zipfile
from urllib.parse import urlsplit

import pymupdf

FONTS_DIR = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
ARIAL = os.path.join(FONTS_DIR, "arial.ttf")
ARIAL_B = os.path.join(FONTS_DIR, "arialbd.ttf")
MM = 72 / 25.4
_office_lock = threading.Lock()


class Erro(ValueError):
    """Erro de uso (mensagem em português para mostrar ao usuário)."""


# ------------------------------------------------------------------ utilitários
def base(nome):
    return re.sub(r"[^\w\-. ()]", "_", os.path.splitext(os.path.basename(nome or "arquivo"))[0])[:80] or "arquivo"


def abrir_pdf(dados, nome="arquivo"):
    try:
        doc = pymupdf.open(stream=dados, filetype="pdf")
    except Exception:
        raise Erro("Não foi possível abrir “%s”: o PDF está corrompido ou não é um PDF." % nome)
    if doc.needs_pass:
        doc.close()
        raise Erro("“%s” tem senha. Use “Desbloquear PDF” primeiro." % nome)
    return doc


def faixa(texto, total, vazio_todas=False):
    """'1-3,5,7-' -> [0,1,2,4,6,...] (índices); levanta Erro se algo estiver fora do documento."""
    texto = (texto or "").strip().lower()
    if not texto:
        if vazio_todas:
            return list(range(total))
        raise Erro("Informe as páginas (por exemplo: 1-3, 5, 8-).")
    out = []
    for tok in re.split(r"[,;\s]+", texto):
        if not tok:
            continue
        m = re.fullmatch(r"(\d+)?\s*-\s*(\d+)?", tok)
        if m and (m.group(1) or m.group(2)):
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else total
        elif tok.isdigit():
            a = b = int(tok)
        elif tok in ("fim", "ultima", "última"):
            a = b = total
        else:
            raise Erro("Páginas inválidas: “%s”." % tok)
        if a < 1 or b > total or a > b:
            raise Erro("As páginas %s estão fora do documento (ele tem %d)." % (tok, total))
        for p in range(a, b + 1):
            if p - 1 not in out:
                out.append(p - 1)
    if not out:
        raise Erro("Nenhuma página escolhida.")
    return out


def salvar(doc, **kw):
    kw.setdefault("garbage", 3)
    kw.setdefault("deflate", True)
    return doc.tobytes(**kw)


def hex_rgb(h, padrao=(0, 0, 0)):
    h = (h or "").lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    try:
        n = int(h, 16)
        return ((n >> 16 & 255) / 255, (n >> 8 & 255) / 255, (n & 255) / 255)
    except Exception:
        return padrao


def num(v, padrao, mini=None, maxi=None):
    try:
        x = float(str(v).replace(",", "."))
    except Exception:
        x = padrao
    if mini is not None:
        x = max(mini, x)
    if maxi is not None:
        x = min(maxi, x)
    return x


def fonte_arial(negrito=False):
    p = ARIAL_B if negrito else ARIAL
    return p if os.path.exists(p) else None


def latin1(txt):
    """True se o texto cabe na fonte padrão do PDF (acentos do português incluídos). Preferida porque o texto
    sai copiável e pesquisável com espaços normais; fora disso usamos o Arial do Windows (Unicode)."""
    try:
        txt.encode("cp1252")
        return True
    except Exception:
        return False


def kb(n):
    return "%d KB" % max(1, round(n / 1024)) if n < 1024 * 1024 else "%.1f MB" % (n / 1024 / 1024)


# ------------------------------------------------------------------ organizar
def juntar(arqs, op):
    out = pymupdf.open()
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        out.insert_pdf(d)
        d.close()
    return [("juntado.pdf", salvar(out))], ["%d arquivos juntados, %d páginas no total." % (len(arqs), out.page_count)]


def dividir(arqs, op):
    modo = op.get("modo", "intervalos")
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        n = d.page_count
        partes = []
        if modo == "cada":
            k = int(num(op.get("cada"), 1, 1, 10000))
            partes = [list(range(i, min(i + k, n))) for i in range(0, n, k)]
        elif modo == "paginas":
            partes = [[i] for i in range(n)]
        else:
            for tok in re.split(r"[;,]+", op.get("intervalos", "")):
                if tok.strip():
                    partes.append(faixa(tok, n))
            if not partes:
                raise Erro("Informe os intervalos (por exemplo: 1-3, 4-6, 7-).")
        for i, pgs in enumerate(partes, 1):
            o = pymupdf.open()
            for p in pgs:
                o.insert_pdf(d, from_page=p, to_page=p)
            saidas.append(("%s-parte-%02d.pdf" % (base(nome), i), salvar(o)))
            o.close()
        d.close()
    return saidas, ["Gerados %d arquivos." % len(saidas)]


def remover(arqs, op):
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        pgs = faixa(op.get("paginas"), d.page_count)
        if len(pgs) >= d.page_count:
            raise Erro("Isso apagaria todas as páginas de “%s”." % nome)
        d.delete_pages(pgs)
        saidas.append((base(nome) + "-sem-paginas.pdf", salvar(d)))
        d.close()
    return saidas, []


def extrair(arqs, op):
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        pgs = faixa(op.get("paginas"), d.page_count)
        d.select(pgs)
        saidas.append((base(nome) + "-extraido.pdf", salvar(d)))
        d.close()
    return saidas, []


def girar(arqs, op):
    ang = int(num(op.get("angulo"), 90, 0, 270))
    if ang not in (90, 180, 270):
        raise Erro("Escolha 90, 180 ou 270 graus.")
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        for i in faixa(op.get("paginas"), d.page_count, vazio_todas=True):
            d[i].set_rotation((d[i].rotation + ang) % 360)
        saidas.append((base(nome) + "-girado.pdf", salvar(d)))
        d.close()
    return saidas, []


def _caixa_conteudo(pg):
    r = pymupdf.Rect()
    for b in pg.get_text("blocks"):
        r |= pymupdf.Rect(b[:4])
    for im in pg.get_image_info():
        r |= pymupdf.Rect(im["bbox"])
    for dr in pg.get_drawings():
        rr = pymupdf.Rect(dr["rect"])
        if rr.width < pg.rect.width * 0.98 or rr.height < pg.rect.height * 0.98:  # ignora fundo da página inteira
            r |= rr
    return r


def cortar(arqs, op):
    auto = bool(op.get("auto"))
    t, b, e, di = (num(op.get(k), 0, 0, 500) * MM for k in ("topo", "base", "esq", "dir"))
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        for pg in d:
            if auto:
                c = _caixa_conteudo(pg)
                if c.is_empty or c.is_infinite:
                    continue
                m = 14
                alvo = pymupdf.Rect(c.x0 - m, c.y0 - m, c.x1 + m, c.y1 + m) & pg.rect
            else:
                alvo = pymupdf.Rect(pg.rect.x0 + e, pg.rect.y0 + t, pg.rect.x1 - di, pg.rect.y1 - b)
            if alvo.width < 20 or alvo.height < 20:
                raise Erro("As margens são grandes demais para a página %d." % (pg.number + 1))
            if pg.rotation:
                alvo = alvo * pg.derotation_matrix
            pg.set_cropbox(alvo)
        saidas.append((base(nome) + "-cortado.pdf", salvar(d)))
        d.close()
    return saidas, ["As páginas foram cortadas (o conteúdo fora da área continua no arquivo, só fica escondido)."]


# ------------------------------------------------------------------ otimizar
def comprimir(arqs, op):
    nivel = op.get("nivel", "media")
    dpi, q = {"baixa": (170, 85), "media": (130, 70), "alta": (96, 50)}.get(nivel, (130, 70))
    saidas, avisos = [], []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        try:
            d.rewrite_images(dpi_threshold=int(dpi * 1.25), dpi_target=dpi, quality=q)
        except Exception:
            pass
        try:
            d.subset_fonts()
        except Exception:
            pass
        novo = salvar(d, garbage=4, clean=True)
        d.close()
        if len(novo) >= len(dados):
            novo = dados
            avisos.append("“%s” já estava bem otimizado: mantive o original." % nome)
        else:
            avisos.append("“%s”: de %s para %s (%d%% menor)." % (nome, kb(len(dados)), kb(len(novo)), round(100 - 100 * len(novo) / len(dados))))
        saidas.append((base(nome) + "-comprimido.pdf", novo))
    return saidas, avisos


def reparar(arqs, op):
    saidas, avisos = [], []
    for nome, dados in arqs:
        try:
            d = pymupdf.open(stream=dados, filetype="pdf")
        except Exception:
            raise Erro("Não consegui recuperar “%s”: o arquivo está muito danificado." % nome)
        if d.needs_pass:
            raise Erro("“%s” tem senha." % nome)
        reparado = d.is_repaired
        if d.page_count == 0:
            raise Erro("“%s” não tem nenhuma página recuperável." % nome)
        saidas.append((base(nome) + "-reparado.pdf", salvar(d, garbage=4, clean=True)))
        avisos.append("“%s”: %s (%d páginas)." % (nome, "estrutura reparada" if reparado else "não havia defeito aparente, reescrito limpo", d.page_count))
        d.close()
    return saidas, avisos


def ocr(arqs, op, ctx):
    saidas, avisos = [], []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        pgs = faixa(op.get("paginas"), d.page_count, vazio_todas=True)
        d.close()
        r = ctx["ocr_pdf"](dados, pgs, op.get("idioma", "por+eng"), bool(op.get("pular", True)))
        import base64
        saidas.append((base(nome) + "-ocr.pdf", base64.b64decode(r["pdf"])))
        avisos.append("“%s”: texto reconhecido em %d página(s). %s" % (nome, r["reconhecidas"], " ".join(r["avisos"])))
    return saidas, avisos


# ------------------------------------------------------------------ editar e segurança
def _centro(pg):
    return pymupdf.Point(pg.rect.width / 2, pg.rect.height / 2)


def marca(arqs, op):
    tipo = op.get("tipo", "texto")
    opac = num(op.get("opacidade"), 30, 1, 100) / 100
    ang = num(op.get("angulo"), 45, -180, 180)
    rep = op.get("posicao") == "repetir"
    cor = hex_rgb(op.get("cor"), (0.6, 0.6, 0.6))
    img_png = None
    if tipo == "imagem":
        from PIL import Image
        import base64
        im = op.get("imagem")
        if not im or not im.get("b64"):
            raise Erro("Escolha a imagem da marca d'água.")
        pil = Image.open(io.BytesIO(base64.b64decode(im["b64"]))).convert("RGBA")
        a = pil.getchannel("A").point(lambda v: int(v * opac))
        pil.putalpha(a)
        if ang:
            pil = pil.rotate(ang, expand=True, resample=Image.BICUBIC)
        buf = io.BytesIO()
        pil.save(buf, "PNG")
        img_png, wh = buf.getvalue(), pil.size
    else:
        texto = (op.get("texto") or "").strip()
        if not texto:
            raise Erro("Digite o texto da marca d'água.")
        fp = None if latin1(texto) else fonte_arial(True)
        tam = num(op.get("tamanho"), 48, 6, 400)
        fonte = pymupdf.Font(fontfile=fp) if fp else pymupdf.Font("hebo")
        larg = fonte.text_length(texto, fontsize=tam)
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        for i in faixa(op.get("paginas"), d.page_count, vazio_todas=True):
            pg = d[i]
            W, H = pg.rect.width, pg.rect.height
            if tipo == "imagem":
                esc = num(op.get("escala"), 50, 5, 300) / 100
                w = wh[0] * esc * 72 / 96
                h = wh[1] * esc * 72 / 96
                if w > W:
                    h, w = h * W / w, W
                alvos = [(W / 2, H / 2)] if not rep else [(x, y) for x in (W * .25, W * .75) for y in (H * .2, H * .5, H * .8)]
                for cx, cy in alvos:
                    pg.insert_image(pymupdf.Rect(cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2), stream=img_png, overlay=True)
            else:
                alvos = [(W / 2, H / 2)] if not rep else [(x, y) for x in (W * .28, W * .72) for y in (H * .18, H * .38, H * .58, H * .78)]
                for cx, cy in alvos:
                    p = pymupdf.Point(cx - larg / 2, cy + tam * 0.3)
                    kw = dict(fontsize=tam, color=cor, fill_opacity=opac, overlay=True, morph=(pymupdf.Point(cx, cy), pymupdf.Matrix(-ang)))
                    if fp:
                        pg.insert_text(p, texto, fontname="WMARCA", fontfile=ARIAL_B, **kw)
                    else:
                        pg.insert_text(p, texto, fontname="hebo", **kw)
        saidas.append((base(nome) + "-marca.pdf", salvar(d)))
        d.close()
    return saidas, []


def numerar(arqs, op):
    pos = op.get("posicao", "baixo-centro")
    fmt = (op.get("formato") or "{n}").strip() or "{n}"
    ini = int(num(op.get("inicio"), 1, 0, 100000))
    tam = num(op.get("tamanho"), 11, 5, 60)
    marg = num(op.get("margem"), 14, 4, 120) * MM
    pular = bool(op.get("pular_primeira"))
    cor = hex_rgb(op.get("cor"), (0.2, 0.2, 0.2))
    fp = None if latin1(fmt) else fonte_arial()
    fonte = pymupdf.Font(fontfile=fp) if fp else pymupdf.Font("helv")
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        total = d.page_count
        alvo = faixa(op.get("paginas"), total, vazio_todas=True)
        for k, i in enumerate(alvo):
            if pular and i == 0:
                continue
            pg = d[i]
            numero = ini + (i - (1 if pular else 0))
            txt = fmt.replace("{n}", str(numero)).replace("{total}", str(total + ini - 1 - (1 if pular else 0)))
            w = fonte.text_length(txt, fontsize=tam)
            R = pg.rect
            if pos.startswith("cima"):
                y = R.y0 + marg + tam
            else:
                y = R.y1 - marg
            if pos.endswith("esquerda"):
                x = R.x0 + marg
            elif pos.endswith("direita"):
                x = R.x1 - marg - w
            else:
                x = (R.x0 + R.x1) / 2 - w / 2
            p = pymupdf.Point(x, y)
            if pg.rotation:
                p = p * pg.derotation_matrix
            kw = dict(fontsize=tam, color=cor, overlay=True, rotate=pg.rotation)
            if fp:
                pg.insert_text(p, txt, fontname="NUMPAG", fontfile=ARIAL, **kw)
            else:
                pg.insert_text(p, txt, fontname="helv", **kw)
        saidas.append((base(nome) + "-numerado.pdf", salvar(d)))
        d.close()
    return saidas, []


def proteger(arqs, op):
    senha = op.get("senha") or ""
    if len(senha) < 4:
        raise Erro("Use uma senha com pelo menos 4 caracteres.")
    perm = pymupdf.PDF_PERM_ACCESSIBILITY
    if op.get("imprimir", True):
        perm |= pymupdf.PDF_PERM_PRINT | pymupdf.PDF_PERM_PRINT_HQ
    if op.get("copiar", True):
        perm |= pymupdf.PDF_PERM_COPY
    if op.get("editar", False):
        perm |= pymupdf.PDF_PERM_MODIFY | pymupdf.PDF_PERM_ANNOTATE | pymupdf.PDF_PERM_FORM | pymupdf.PDF_PERM_ASSEMBLE
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        saidas.append((base(nome) + "-protegido.pdf", d.tobytes(garbage=3, deflate=True, encryption=pymupdf.PDF_ENCRYPT_AES_256,
                                                                owner_pw=senha, user_pw=senha, permissions=int(perm))))
        d.close()
    return saidas, ["Para abrir, será pedida a senha que você digitou. Guarde-a: não há como recuperar."]


def desbloquear(arqs, op):
    senha = op.get("senha") or ""
    saidas, avisos = [], []
    for nome, dados in arqs:
        try:
            d = pymupdf.open(stream=dados, filetype="pdf")
        except Exception:
            raise Erro("Não foi possível abrir “%s”." % nome)
        if d.needs_pass:
            if not senha or not d.authenticate(senha):
                raise Erro("A senha não abre “%s”. Esta ferramenta só remove a proteção de quem sabe a senha." % nome)
            avisos.append("“%s”: senha aceita, proteção removida." % nome)
        else:
            avisos.append("“%s” abre sem senha: removi as restrições de impressão, cópia e edição." % nome)
        saidas.append((base(nome) + "-desbloqueado.pdf", d.tobytes(garbage=3, deflate=True, encryption=pymupdf.PDF_ENCRYPT_NONE)))
        d.close()
    return saidas, avisos


def comparar(arqs, op):
    if len(arqs) != 2:
        raise Erro("Escolha exatamente 2 PDFs para comparar.")
    (na, da), (nb, db) = arqs
    A, B = abrir_pdf(da, na), abrir_pdf(db, nb)
    rel = pymupdf.open()
    fp = fonte_arial()
    fpb = fonte_arial(True)
    estado = {"pg": None, "y": 0}

    def nova():
        estado["pg"] = rel.new_page(width=595, height=842)
        estado["y"] = 50

    def linha(txt, cor=(0, 0, 0), tam=9, negrito=False):
        for parte in textwrap.wrap(txt, 105) or [""]:
            if estado["pg"] is None or estado["y"] > 800:
                nova()
            kw = dict(fontsize=tam, color=cor)
            f = None if latin1(parte) else (fpb if negrito else fp)
            if f:
                estado["pg"].insert_text((40, estado["y"]), parte, fontname="RELB" if negrito else "REL", fontfile=f, **kw)
            else:
                estado["pg"].insert_text((40, estado["y"]), parte, fontname="hebo" if negrito else "helv", **kw)
            estado["y"] += tam * 1.5

    linha("Comparação de PDFs", tam=16, negrito=True)
    linha("A (antigo): %s  [%d páginas]" % (na, A.page_count))
    linha("B (novo):   %s  [%d páginas]" % (nb, B.page_count))
    linha("Em vermelho: o que só existe em A. Em verde: o que só existe em B.", cor=(.35, .35, .35))
    estado["y"] += 8
    com_dif = 0
    for i in range(max(A.page_count, B.page_count)):
        ta = A[i].get_text().splitlines() if i < A.page_count else []
        tb = B[i].get_text().splitlines() if i < B.page_count else []
        ta = [l.strip() for l in ta if l.strip()]
        tb = [l.strip() for l in tb if l.strip()]
        difs = [l for l in difflib.ndiff(ta, tb) if l[:1] in "+-"]
        if not difs:
            continue
        com_dif += 1
        linha("Página %d" % (i + 1), negrito=True, tam=11)
        for l in difs:
            linha(("- " if l[0] == "-" else "+ ") + l[2:], cor=(0.75, 0.1, 0.1) if l[0] == "-" else (0.05, 0.5, 0.2))
        estado["y"] += 6
    if com_dif == 0:
        linha("Nenhuma diferença de texto encontrada.", cor=(0.05, 0.5, 0.2), negrito=True, tam=12)
    A.close(); B.close()
    resumo = ["Diferenças de texto em %d página(s)." % com_dif if com_dif else "Os textos são iguais. Obs.: só o texto é comparado, não a aparência."]
    return [("comparacao.pdf", salvar(rel))], resumo


# ------------------------------------------------------------------ converter para PDF
def img2pdf(arqs, op):
    from PIL import Image, ImageOps
    tamanho = op.get("tamanho", "a4")
    orient = op.get("orientacao", "auto")
    marg = num(op.get("margem"), 10 if tamanho != "imagem" else 0, 0, 100) * MM
    sep = bool(op.get("separado"))
    out = pymupdf.open()
    saidas = []
    for nome, dados in arqs:
        try:
            im = Image.open(io.BytesIO(dados))
            im = ImageOps.exif_transpose(im)
        except Exception:
            raise Erro("“%s” não é uma imagem válida (use JPG, PNG, WEBP, BMP, TIFF ou GIF)." % nome)
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA")
            fundo = Image.new("RGB", im.size, (255, 255, 255))
            fundo.paste(im, mask=im.getchannel("A"))
            im = fundo
        else:
            im = im.convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=92)
        jpg = buf.getvalue()
        iw, ih = im.size
        if tamanho == "imagem":
            pw, ph = iw * 72 / 96, ih * 72 / 96
        else:
            pw, ph = {"a4": (595.28, 841.89), "carta": (612, 792)}.get(tamanho, (595.28, 841.89))
            paisagem = (orient == "paisagem") or (orient == "auto" and iw > ih)
            if paisagem:
                pw, ph = ph, pw
        alvo = out if not sep else pymupdf.open()
        pg = alvo.new_page(width=pw, height=ph)
        area = pymupdf.Rect(marg, marg, pw - marg, ph - marg)
        esc = min(area.width / iw, area.height / ih)
        w, h = iw * esc, ih * esc
        r = pymupdf.Rect(area.x0 + (area.width - w) / 2, area.y0 + (area.height - h) / 2, 0, 0)
        r.x1, r.y1 = r.x0 + w, r.y0 + h
        pg.insert_image(r, stream=jpg)
        if sep:
            saidas.append((base(nome) + ".pdf", salvar(alvo)))
            alvo.close()
    if not sep:
        saidas = [("imagens.pdf", salvar(out))]
    return saidas, []


_PS_OFFICE = r'''
param([string]$Entrada, [string]$Saida, [string]$Tipo)
$ErrorActionPreference = 'Stop'
$app = $null
try {
  if ($Tipo -eq 'word') {
    $app = New-Object -ComObject Word.Application
    $app.Visible = $false; $app.DisplayAlerts = 0; $app.AutomationSecurity = 3
    $doc = $app.Documents.Open($Entrada, $false, $true, $false)
    $doc.ExportAsFixedFormat($Saida, 17)
    $doc.Close(0)
  } elseif ($Tipo -eq 'excel') {
    $app = New-Object -ComObject Excel.Application
    $app.Visible = $false; $app.DisplayAlerts = $false; $app.AutomationSecurity = 3
    $wb = $app.Workbooks.Open($Entrada, 0, $true)
    $wb.ExportAsFixedFormat(0, $Saida)
    $wb.Close($false)
  } elseif ($Tipo -eq 'powerpoint') {
    $app = New-Object -ComObject PowerPoint.Application
    $app.AutomationSecurity = 3
    $pr = $app.Presentations.Open($Entrada, -1, 0, 0)
    $pr.SaveAs($Saida, 32)
    $pr.Close()
  }
} finally {
  if ($app) { try { $app.Quit() } catch { } }
}
'''
_OFFICE = {".doc": "word", ".docx": "word", ".rtf": "word", ".odt": "word", ".xls": "excel", ".xlsx": "excel", ".xlsm": "excel",
           ".csv": "excel", ".ppt": "powerpoint", ".pptx": "powerpoint", ".pps": "powerpoint", ".ppsx": "powerpoint"}


def office2pdf(arqs, op):
    saidas, avisos = [], []
    with _office_lock, tempfile.TemporaryDirectory() as tmp:
        ps1 = os.path.join(tmp, "converter.ps1")
        with open(ps1, "w", encoding="utf-8-sig") as f:
            f.write(_PS_OFFICE)
        for nome, dados in arqs:
            ext = os.path.splitext(nome)[1].lower()
            tipo = _OFFICE.get(ext)
            if not tipo:
                raise Erro("“%s”: formato não aceito (use Word, Excel ou PowerPoint)." % nome)
            ent = os.path.join(tmp, uuid.uuid4().hex + ext)
            sai = os.path.join(tmp, uuid.uuid4().hex + ".pdf")
            with open(ent, "wb") as f:
                f.write(dados)
            try:
                r = subprocess.run(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", ps1, "-Entrada", ent, "-Saida", sai, "-Tipo", tipo],
                                   capture_output=True, text=True, timeout=240, creationflags=0x08000000)
            except subprocess.TimeoutExpired:
                raise Erro("O Office demorou demais para converter “%s”." % nome)
            if r.returncode != 0 or not os.path.exists(sai):
                detalhe = (r.stderr or r.stdout or "").strip().splitlines()
                raise Erro("O Office não conseguiu converter “%s”%s" % (nome, (": " + detalhe[0][:160]) if detalhe else ". Ele está instalado?"))
            with open(sai, "rb") as f:
                saidas.append((base(nome) + ".pdf", f.read()))
    avisos.append("Convertido pelo Microsoft Office deste computador, com as macros desativadas.")
    return saidas, avisos


def _url_publica(url):
    p = urlsplit(url)
    if p.scheme not in ("http", "https") or not p.hostname:
        raise Erro("Digite um endereço começando com http:// ou https://")
    try:
        infos = socket.getaddrinfo(p.hostname, p.port or (443 if p.scheme == "https" else 80))
    except Exception:
        raise Erro("Não consegui encontrar o endereço “%s”." % p.hostname)
    for i in infos:
        ip = ipaddress.ip_address(i[4][0])
        if not ip.is_global:
            raise Erro("Por segurança, só páginas públicas da internet são aceitas (não endereços da rede local).")
    return url


def html2pdf(arqs, op):
    edge = next((p for p in (os.path.join(os.environ.get("ProgramFiles(x86)", ""), r"Microsoft\Edge\Application\msedge.exe"),
                             os.path.join(os.environ.get("ProgramFiles", ""), r"Microsoft\Edge\Application\msedge.exe"),
                             os.path.join(os.environ.get("ProgramFiles", ""), r"Google\Chrome\Application\chrome.exe")) if os.path.exists(p)), None)
    if not edge:
        raise Erro("Não encontrei o Edge nem o Chrome neste computador.")
    fonte = op.get("fonte", "url")
    saidas = []
    with tempfile.TemporaryDirectory() as tmp:
        alvos = []
        if fonte == "url":
            url = _url_publica((op.get("url") or "").strip())
            alvos.append(("pagina", url))
        else:
            if not arqs:
                raise Erro("Escolha um arquivo .html.")
            for nome, dados in arqs:
                txt = dados.decode("utf-8", "replace")
                txt = "<style>@page{size:A4;margin:12mm}</style>" + txt
                caminho = os.path.join(tmp, uuid.uuid4().hex + ".html")
                with open(caminho, "w", encoding="utf-8") as f:
                    f.write(txt)
                alvos.append((base(nome), "file:///" + caminho.replace("\\", "/")))
        for nome, url in alvos:
            sai = os.path.join(tmp, uuid.uuid4().hex + ".pdf")
            perfil = os.path.join(tmp, "perfil-" + uuid.uuid4().hex[:6])
            try:
                subprocess.run([edge, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--user-data-dir=" + perfil,
                                "--print-to-pdf=" + sai, url], capture_output=True, timeout=90, creationflags=0x08000000)
            except subprocess.TimeoutExpired:
                raise Erro("A página demorou demais para carregar.")
            if not os.path.exists(sai):
                raise Erro("Não foi possível gerar o PDF desta página.")
            with open(sai, "rb") as f:
                saidas.append((nome + ".pdf", f.read()))
    return saidas, ["Se a página usa muitos recursos externos, o resultado pode diferir do que você vê no navegador."]


# ------------------------------------------------------------------ converter de PDF
def pdf2img(arqs, op):
    fmt = op.get("formato", "jpg")
    dpi = int(num(op.get("dpi"), 150, 50, 400))
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        pgs = faixa(op.get("paginas"), d.page_count, vazio_todas=True)
        for i in pgs:
            pix = d[i].get_pixmap(dpi=dpi, alpha=False)
            b = pix.tobytes("jpg", jpg_quality=90) if fmt == "jpg" else pix.tobytes("png")
            saidas.append(("%s-pagina-%03d.%s" % (base(nome), i + 1, fmt), b))
        d.close()
    return saidas, ["%d imagem(ns) geradas a %d dpi." % (len(saidas), dpi)]


def pdf2word(arqs, op, ctx):
    import base64
    saidas, avisos = [], []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        pgs = faixa(op.get("paginas"), d.page_count, vazio_todas=True)
        if len(pgs) < d.page_count:
            d.select(pgs)
            dados = d.tobytes()
        d.close()
        r = ctx["converter_docx"](dados, 0, None, "auto", op.get("idioma", "por+eng"))
        saidas.append((base(nome) + ".docx", base64.b64decode(r["docx"])))
        avisos += ["“%s”: %s" % (nome, a) for a in r["avisos"]]
    return saidas, avisos


_NUM_BR = re.compile(r"^-?\d{1,3}(\.\d{3})*(,\d+)?$|^-?\d+(,\d+)?$")


def _valor(s, converter):
    s = (s or "").strip()
    if converter and s and _NUM_BR.match(s):
        try:
            v = float(s.replace(".", "").replace(",", "."))
            return int(v) if v.is_integer() and "," not in s else v
        except Exception:
            return s
    return s


def pdf2xlsx(arqs, op):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    conv = bool(op.get("numeros", True))
    saidas, avisos = [], []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        wb = Workbook()
        wb.remove(wb.active)
        n_tab = 0
        for i in faixa(op.get("paginas"), d.page_count, vazio_todas=True):
            try:
                tabs = d[i].find_tables().tables
            except Exception:
                tabs = []
            for k, t in enumerate(tabs, 1):
                linhas = t.extract()
                if not linhas or all(not any(c for c in l) for l in linhas):
                    continue
                n_tab += 1
                ws = wb.create_sheet(("P%d-T%d" % (i + 1, k))[:31])
                for r, l in enumerate(linhas, 1):
                    for c, v in enumerate(l, 1):
                        ws.cell(r, c, _valor((v or "").replace("\n", " "), conv and r > 1))
                for c in ws[1]:
                    c.font = Font(bold=True)
                for c in range(1, ws.max_column + 1):
                    m = max(len(str(ws.cell(r, c).value or "")) for r in range(1, ws.max_row + 1))
                    ws.column_dimensions[get_column_letter(c)].width = min(60, max(8, m + 2))
        if n_tab == 0:
            ws = wb.create_sheet("Texto")
            for i in faixa(op.get("paginas"), d.page_count, vazio_todas=True):
                for l in d[i].get_text().splitlines():
                    if l.strip():
                        ws.append([l.strip()])
            avisos.append("“%s”: não encontrei tabelas; coloquei o texto, uma linha por célula." % nome)
        else:
            avisos.append("“%s”: %d tabela(s) encontrada(s). Confira os valores." % (nome, n_tab))
        buf = io.BytesIO()
        wb.save(buf)
        saidas.append((base(nome) + ".xlsx", buf.getvalue()))
        d.close()
    return saidas, avisos


def pdf2pptx(arqs, op):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    modo = op.get("modo", "imagem")
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        prs = Presentation()
        prs.slide_width = Emu(int(d[0].rect.width * 12700))
        prs.slide_height = Emu(int(d[0].rect.height * 12700))
        branco = prs.slide_layouts[6]
        for i in faixa(op.get("paginas"), d.page_count, vazio_todas=True):
            pg = d[i]
            sl = prs.slides.add_slide(branco)
            if modo == "imagem":
                pix = pg.get_pixmap(dpi=150, alpha=False)
                sl.shapes.add_picture(io.BytesIO(pix.tobytes("jpg", jpg_quality=88)), 0, 0, prs.slide_width, prs.slide_height)
            else:
                for b in pg.get_text("dict")["blocks"]:
                    if b.get("type") == 1:
                        continue
                    if b.get("type") != 0:
                        continue
                    texto = "\n".join("".join(s["text"] for s in l["spans"]) for l in b["lines"]).strip()
                    if not texto:
                        continue
                    s0 = b["lines"][0]["spans"][0]
                    x0, y0, x1, y1 = b["bbox"]
                    tb = sl.shapes.add_textbox(Emu(int(x0 * 12700)), Emu(int(y0 * 12700)), Emu(int((x1 - x0 + 6) * 12700)), Emu(int((y1 - y0) * 12700)))
                    tf = tb.text_frame
                    tf.word_wrap = True
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    run = tf.paragraphs[0].add_run()
                    run.text = texto
                    run.font.size = Pt(max(6, s0["size"]))
                    run.font.bold = bool(s0["flags"] & 16)
                    run.font.italic = bool(s0["flags"] & 2)
                    c = s0["color"]
                    run.font.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
        buf = io.BytesIO()
        prs.save(buf)
        saidas.append((base(nome) + ".pptx", buf.getvalue()))
        d.close()
    return saidas, ["Cada página virou um slide idêntico ao PDF, mas não editável." if modo == "imagem"
                    else "Caixas de texto editáveis nas mesmas posições; imagens e desenhos do PDF não foram levados."]


def pdf2txt(arqs, op):
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        partes = [d[i].get_text() for i in faixa(op.get("paginas"), d.page_count, vazio_todas=True)]
        texto = "\n\n".join(partes)
        saidas.append((base(nome) + ".txt", texto.encode("utf-8")))
        d.close()
    av = [] if any(s[1].strip() for s in saidas) else ["Nenhum texto encontrado: o PDF parece escaneado. Use o OCR antes."]
    return saidas, av


# ------------------------------------------------------------------ cabeçalho e rodapé
def _img_op(op, chave, rotulo):
    """Imagem enviada como opção ({nome, b64}) -> bytes PNG/JPG, ou None."""
    import base64
    im = op.get(chave)
    if not im or not isinstance(im, dict) or not im.get("b64"):
        return None
    try:
        return base64.b64decode(im["b64"])
    except Exception:
        raise Erro("%s inválido." % rotulo)


def cabecalho(arqs, op):
    import datetime
    campos = {k: (op.get(k) or "").strip() for k in ("cab_esq", "cab_centro", "cab_dir", "rod_esq", "rod_centro", "rod_dir")}
    logo = _img_op(op, "logo", "Logotipo")
    if not any(campos.values()) and not logo:
        raise Erro("Escreva algum texto para o cabeçalho ou o rodapé (ou escolha um logotipo).")
    tam = num(op.get("tamanho"), 9, 5, 40)
    marg = num(op.get("margem"), 10, 3, 60) * MM
    cor = hex_rgb(op.get("cor"), (0.3, 0.3, 0.3))
    pular = bool(op.get("pular_primeira"))
    linha = bool(op.get("linha"))
    logo_h = num(op.get("logo_altura"), 12, 4, 60) * MM
    logo_pos = op.get("logo_pos", "cab_esq")
    hoje = datetime.date.today().strftime("%d/%m/%Y")
    tudo = "".join(campos.values())
    fp = None if latin1(tudo) else fonte_arial()
    fonte = pymupdf.Font(fontfile=fp) if fp else pymupdf.Font("helv")
    if logo:
        try:
            lw, lh = pymupdf.Pixmap(logo).width, pymupdf.Pixmap(logo).height
        except Exception:
            raise Erro("O logotipo não é uma imagem válida (use PNG ou JPG).")
    saidas = []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        total = d.page_count
        for i in faixa(op.get("paginas"), total, vazio_todas=True):
            if pular and i == 0:
                continue
            pg = d[i]
            R = pg.rect  # já considera a rotação (coordenadas "como se vê")
            subs = lambda t: (t.replace("{n}", str(i + 1)).replace("{total}", str(total)).replace("{data}", hoje)
                              .replace("{arquivo}", os.path.splitext(os.path.basename(nome))[0]))
            for chave, t in campos.items():
                if not t:
                    continue
                t = subs(t)
                w = fonte.text_length(t, fontsize=tam)
                y = R.y0 + marg + tam if chave.startswith("cab") else R.y1 - marg
                x = R.x0 + marg if chave.endswith("esq") else (R.x1 - marg - w if chave.endswith("dir") else (R.x0 + R.x1) / 2 - w / 2)
                p = pymupdf.Point(x, y) * pg.derotation_matrix
                kw = dict(fontsize=tam, color=cor, overlay=True, rotate=pg.rotation)
                if fp:
                    pg.insert_text(p, t, fontname="CABROD", fontfile=fp, **kw)
                else:
                    pg.insert_text(p, t, fontname="helv", **kw)
            if linha:
                for yy in ([R.y0 + marg + tam * 1.6] if any(campos[k] for k in campos if k.startswith("cab")) or (logo and logo_pos.startswith("cab")) else []) + \
                          ([R.y1 - marg - tam * 1.4] if any(campos[k] for k in campos if k.startswith("rod")) else []):
                    a = pymupdf.Point(R.x0 + marg, yy) * pg.derotation_matrix
                    b = pymupdf.Point(R.x1 - marg, yy) * pg.derotation_matrix
                    pg.draw_line(a, b, color=cor, width=0.5, overlay=True)
            if logo:
                h = logo_h
                w = h * lw / lh
                y0 = R.y0 + marg - 2 if logo_pos.startswith("cab") else R.y1 - marg - h + 2
                x0 = R.x0 + marg if logo_pos.endswith("esq") else (R.x1 - marg - w if logo_pos.endswith("dir") else (R.x0 + R.x1) / 2 - w / 2)
                rr = pymupdf.Rect(x0, y0, x0 + w, y0 + h) * pg.derotation_matrix
                pg.insert_image(rr, stream=logo, overlay=True, rotate=pg.rotation)
        saidas.append((base(nome) + "-cabecalho.pdf", salvar(d)))
        d.close()
    return saidas, []


# ------------------------------------------------------------------ busca de texto (tarja automática, buscar e substituir)
PADROES = {
    "cpf": ("CPF", r"(?<!\d)\d{3}\.?\d{3}\.?\d{3}\s?-?\s?\d{2}(?!\d)"),
    "cnpj": ("CNPJ", r"(?<!\d)\d{2}\.?\d{3}\.?\d{3}\s?/?\s?\d{4}\s?-?\s?\d{2}(?!\d)"),
    "rg": ("RG", r"(?<![\d.])\d{1,2}\.\d{3}\.\d{3}-?[\dXx](?![\d.])"),
    "email": ("E-mail", r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"),
    "telefone": ("Telefone", r"(?<!\d)(?:\+?55\s?)?(?:\(?\d{2}\)?\s?)?9?\s?\d{4}\s?-?\s?\d{4}(?!\d)"),
    "cep": ("CEP", r"(?<!\d)\d{5}-\d{3}(?!\d)"),
    "placa": ("Placa", r"\b[A-Z]{3}-?\d[A-Z0-9]\d{2}\b"),
    "data": ("Data", r"(?<!\d)\d{2}/\d{2}/\d{4}(?!\d)"),
}


def _regexes(op):
    """Monta a lista [(rótulo, regex)] a partir das opções: padrões marcados + termos digitados."""
    out = []
    for k in op.get("padroes") or []:
        if k in PADROES:
            out.append((PADROES[k][0], re.compile(PADROES[k][1])))
    termos = op.get("termos")
    if isinstance(termos, str):
        termos = [t for t in re.split(r"[\n;]+", termos)]
    flags = 0 if op.get("maiusculas") else re.IGNORECASE
    for t in termos or []:
        t = (t or "").strip()
        if not t:
            continue
        if len(t) > 300:
            raise Erro("Termo de busca grande demais.")
        pad = re.escape(t)
        if op.get("inteira"):
            pad = r"(?<!\w)" + pad + r"(?!\w)"
        out.append(("Termo", re.compile(pad, flags)))
    if not out:
        raise Erro("Marque o que procurar ou digite uma palavra.")
    return out


def ocorrencias(pg, regs):
    """Procura as expressões em cada linha da página; devolve [(rótulo, texto, Rect)] com a área exata dos caracteres."""
    achados = []
    raw = pg.get_text("rawdict", flags=0)
    for b in raw["blocks"]:
        for l in b.get("lines", []):
            chars = [c for s in l["spans"] for c in s["chars"]]
            if not chars:
                continue
            texto = "".join(c["c"] for c in chars)
            for rot, rg in regs:
                for m in rg.finditer(texto):
                    if m.end() <= m.start():
                        continue
                    r = pymupdf.Rect()
                    for c in chars[m.start():m.end()]:
                        r |= pymupdf.Rect(c["bbox"])
                    if not r.is_empty:
                        achados.append((rot, m.group(0), r))
    # padrões que se sobrepõem (ex.: CPF dentro de "Termo"): mantém um só por área
    unicos = []
    for a in achados:
        if not any(abs(a[2].x0 - u[2].x0) < 0.5 and abs(a[2].y0 - u[2].y0) < 0.5 and abs(a[2].x1 - u[2].x1) < 0.5 for u in unicos):
            unicos.append(a)
    return unicos


def buscar(dados, op):
    """Para a pré-visualização no editor: [{pagina, rotulo, texto, rect}]."""
    d = abrir_pdf(dados)
    regs = _regexes(op)
    res = []
    try:
        for pg in d:
            for rot, txt, r in ocorrencias(pg, regs):
                res.append({"pagina": pg.number, "rotulo": rot, "texto": txt, "rect": [round(v, 2) for v in r]})
                if len(res) >= 5000:
                    return res
    finally:
        d.close()
    return res


def tarja_auto(arqs, op):
    regs = _regexes(op)
    cor = hex_rgb(op.get("cor"), (0, 0, 0))
    saidas, avisos = [], []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        cont = {}
        sem_texto = 0
        for i in faixa(op.get("paginas"), d.page_count, vazio_todas=True):
            pg = d[i]
            if not pg.get_text().strip():
                sem_texto += 1
            achados = ocorrencias(pg, regs)
            for rot, _t, r in achados:
                pg.add_redact_annot(r + (-0.5, -0.5, 0.5, 0.5), fill=cor)
                cont[rot] = cont.get(rot, 0) + 1
            if achados:
                pg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                                    text=pymupdf.PDF_REDACT_TEXT_REMOVE)
        if op.get("metadados", True):
            d.set_metadata({})
            try:
                d.del_xml_metadata()
            except Exception:
                pass
        total = sum(cont.values())
        saidas.append((base(nome) + "-tarjado.pdf", salvar(d, garbage=4, clean=True)))
        d.close()
        if total:
            avisos.append("“%s”: %d trecho(s) apagado(s) de verdade (%s)." % (nome, total, ", ".join("%s: %d" % kv for kv in cont.items())))
        else:
            avisos.append("“%s”: nada encontrado." % nome)
        if sem_texto:
            avisos.append("“%s”: %d página(s) sem texto (escaneadas). Use o OCR antes, senão os dados nelas não são achados." % (nome, sem_texto))
    return saidas, avisos


def substituir(arqs, op, ctx):
    busca = op.get("busca") or ""
    if not busca.strip():
        raise Erro("Digite o texto a procurar.")
    troca = op.get("troca") or ""
    saidas, avisos = [], []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        pgs = faixa(op.get("paginas"), d.page_count, vazio_todas=True)
        d.close()
        r = ctx["substituir_texto"](dados, busca, troca, bool(op.get("maiusculas")), bool(op.get("inteira")), pgs)
        import base64
        saidas.append((base(nome) + "-substituido.pdf", base64.b64decode(r["pdf"])))
        avisos.append("“%s”: %d ocorrência(s) trocada(s)." % (nome, r["trocas"]) if r["trocas"] else "“%s”: “%s” não foi encontrado." % (nome, busca))
        avisos += r["avisos"]
    return saidas, avisos


# ------------------------------------------------------------------ imagens (estilo iLoveIMG)
_FMT_PIL = {"jpg": "JPEG", "jpeg": "JPEG", "png": "PNG", "webp": "WEBP", "bmp": "BMP", "gif": "GIF", "tif": "TIFF", "tiff": "TIFF"}


def _abrir_img(nome, dados):
    from PIL import Image, ImageOps
    try:
        im = Image.open(io.BytesIO(dados))
        im.load()
        return ImageOps.exif_transpose(im)
    except Exception:
        raise Erro("“%s” não é uma imagem válida." % nome)


def _salvar_img(im, fmt, qualidade=85):
    """Salva no formato pedido (jpg/png/webp/...) cuidando da transparência; devolve (bytes, extensão)."""
    from PIL import Image
    fmt = fmt.lower().replace("jpeg", "jpg")
    pil = _FMT_PIL.get(fmt, "PNG")
    if pil in ("JPEG", "BMP") and im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        fundo = Image.new("RGB", im.size, (255, 255, 255))
        fundo.paste(im, mask=im.getchannel("A"))
        im = fundo
    elif pil == "JPEG" and im.mode not in ("RGB", "L"):
        im = im.convert("RGB")
    buf = io.BytesIO()
    kw = {}
    if pil == "JPEG":
        kw = dict(quality=int(qualidade), optimize=True, progressive=True)
    elif pil == "WEBP":
        kw = dict(quality=int(qualidade), method=6)
    elif pil == "PNG":
        kw = dict(optimize=True)
    im.save(buf, pil, **kw)
    return buf.getvalue(), ("jpg" if pil == "JPEG" else fmt)


def _fmt_de(nome):
    e = os.path.splitext(nome)[1].lower().lstrip(".")
    return {"jpeg": "jpg", "tif": "tiff"}.get(e, e) or "png"


def img_comprimir(arqs, op):
    q = {"baixa": 85, "media": 72, "alta": 55}.get(op.get("nivel", "media"), 72)
    maxlado = int(num(op.get("max_lado"), 0, 0, 20000))
    saidas, avisos = [], []
    for nome, dados in arqs:
        im = _abrir_img(nome, dados)
        if maxlado and max(im.size) > maxlado:
            im.thumbnail((maxlado, maxlado))
        fmt = _fmt_de(nome)
        if fmt == "png" and op.get("nivel") == "alta":
            # PNG: reduz para 256 cores (fica bem menor; ótimo para prints e desenhos)
            im = im.convert("RGBA").quantize(256, method=2) if im.mode in ("RGBA", "LA", "P") else im.convert("RGB").quantize(256)
        if fmt in ("bmp", "tiff", "gif"):
            fmt = "jpg" if im.mode not in ("RGBA", "LA", "P") else "png"
        novo, ext = _salvar_img(im, fmt, q)
        if len(novo) >= len(dados) and not maxlado:
            novo, ext = dados, _fmt_de(nome)
            avisos.append("“%s” já estava otimizada: mantive a original." % nome)
        else:
            avisos.append("“%s”: de %s para %s (%d%% menor)." % (nome, kb(len(dados)), kb(len(novo)), max(0, round(100 - 100 * len(novo) / len(dados)))))
        saidas.append(("%s-comprimida.%s" % (base(nome), ext), novo))
    return saidas, avisos


def img_redimensionar(arqs, op):
    from PIL import Image
    modo = op.get("modo", "porcentagem")
    nao_aumentar = op.get("nao_aumentar", True)
    saidas = []
    for nome, dados in arqs:
        im = _abrir_img(nome, dados)
        W, H = im.size
        if modo == "porcentagem":
            p = num(op.get("porcentagem"), 50, 1, 1000) / 100
            w, h = W * p, H * p
        else:
            lw, lh = int(num(op.get("largura"), 0, 0, 20000)), int(num(op.get("altura"), 0, 0, 20000))
            if not lw and not lh:
                raise Erro("Informe a largura ou a altura em pixels.")
            if op.get("proporcao", True):
                esc = min(lw / W if lw else 1e9, lh / H if lh else 1e9)
                w, h = W * esc, H * esc
            else:
                w, h = lw or W, lh or H
        if nao_aumentar and (w > W or h > H):
            esc = min(W / w, H / h)
            w, h = w * esc, h * esc
        w, h = max(1, round(w)), max(1, round(h))
        im2 = im.resize((w, h), Image.LANCZOS)
        b, ext = _salvar_img(im2, _fmt_de(nome), 90)
        saidas.append(("%s-%dx%d.%s" % (base(nome), w, h, ext), b))
    return saidas, []


def img_recortar(arqs, op):
    modo = op.get("modo", "proporcao")
    saidas = []
    for nome, dados in arqs:
        im = _abrir_img(nome, dados)
        W, H = im.size
        if modo == "margens":
            t, b_, e, di = (int(num(op.get(k), 0, 0, 20000)) for k in ("topo", "base", "esq", "dir"))
            caixa = (e, t, W - di, H - b_)
        else:
            a, b2 = {"1:1": (1, 1), "4:3": (4, 3), "3:4": (3, 4), "16:9": (16, 9), "9:16": (9, 16), "3:2": (3, 2), "2:3": (2, 3), "a4": (210, 297)}.get(op.get("proporcao", "1:1"), (1, 1))
            if W / H > a / b2:
                nw = round(H * a / b2)
                x = (W - nw) // 2
                caixa = (x, 0, x + nw, H)
            else:
                nh = round(W * b2 / a)
                y = (H - nh) // 2
                caixa = (0, y, W, y + nh)
        if caixa[2] - caixa[0] < 2 or caixa[3] - caixa[1] < 2:
            raise Erro("O recorte de “%s” ficou vazio: as margens são maiores que a imagem (%dx%d)." % (nome, W, H))
        b, ext = _salvar_img(im.crop(caixa), _fmt_de(nome), 92)
        saidas.append(("%s-recortada.%s" % (base(nome), ext), b))
    return saidas, []


def img_converter(arqs, op):
    fmt = op.get("formato", "jpg")
    if fmt not in ("jpg", "png", "webp", "bmp", "gif", "tiff"):
        raise Erro("Formato não suportado.")
    q = int(num(op.get("qualidade"), 90, 10, 100))
    saidas = []
    for nome, dados in arqs:
        im = _abrir_img(nome, dados)
        b, ext = _salvar_img(im, fmt, q)
        saidas.append(("%s.%s" % (base(nome), ext), b))
    return saidas, []


def img_girar(arqs, op):
    from PIL import Image, ImageOps
    ang = int(num(op.get("angulo"), 90, 0, 270))
    saidas = []
    for nome, dados in arqs:
        im = _abrir_img(nome, dados)
        if ang:
            im = im.rotate(-ang, expand=True)
        if op.get("espelhar") == "horizontal":
            im = ImageOps.mirror(im)
        elif op.get("espelhar") == "vertical":
            im = ImageOps.flip(im)
        b, ext = _salvar_img(im, _fmt_de(nome), 92)
        saidas.append(("%s-girada.%s" % (base(nome), ext), b))
    return saidas, []


# ------------------------------------------------------------------ PDF/A (arquivamento)
def ghostscript():
    import glob
    import sys
    cands = [os.environ.get("EDITOR_PDF_GS") or ""]
    meipass = getattr(sys, "_MEIPASS", None)
    if meipass:
        cands.append(os.path.join(meipass, "gs", "bin", "gswin64c.exe"))
    cands += sorted(glob.glob(os.path.join(os.environ.get("ProgramFiles", r"C:\Program Files"), "gs", "gs*", "bin", "gswin64c.exe")), reverse=True)
    for c in cands:
        if c and os.path.exists(c):
            return c
    return None


def pdfa(arqs, op):
    gs = ghostscript()
    if not gs:
        raise Erro("O conversor de PDF/A (Ghostscript) não está instalado neste computador.")
    raiz = os.path.dirname(os.path.dirname(gs))
    icc = os.path.join(raiz, "iccprofiles", "srgb.icc")
    if not os.path.exists(icc):
        raise Erro("Falta o perfil de cores do Ghostscript (iccprofiles\\srgb.icc).")
    nivel = op.get("nivel", "2")
    if nivel not in ("1", "2", "3"):
        nivel = "2"
    saidas, avisos = [], []
    with tempfile.TemporaryDirectory() as tmp:
        defps = os.path.join(tmp, "pdfa_def.ps")
        with open(defps, "w", encoding="latin-1") as f:
            f.write("[ /Title (Documento) /DOCINFO pdfmark\n"
                    "/ICCProfile (%s) def\n"
                    "[/_objdef {icc_PDFA} /type /stream /OBJ pdfmark\n"
                    "[{icc_PDFA} << /N 3 >> /PUT pdfmark\n"
                    "[{icc_PDFA} ICCProfile (r) file /PUT pdfmark\n"
                    "[/_objdef {OutputIntent_PDFA} /type /dict /OBJ pdfmark\n"
                    "[{OutputIntent_PDFA} << /Type /OutputIntent /S /GTS_PDFA1 /DestOutputProfile {icc_PDFA} "
                    "/OutputConditionIdentifier (sRGB) >> /PUT pdfmark\n"
                    "[{Catalog} << /OutputIntents [ {OutputIntent_PDFA} ] >> /PUT pdfmark\n" % icc.replace("\\", "/"))
        for nome, dados in arqs:
            d = abrir_pdf(dados, nome)
            titulo = (d.metadata or {}).get("title") or base(nome)
            d.close()
            ent = os.path.join(tmp, uuid.uuid4().hex + ".pdf")
            sai = os.path.join(tmp, uuid.uuid4().hex + ".pdf")
            with open(ent, "wb") as f:
                f.write(dados)
            cmd = [gs, "-dPDFA=%s" % nivel, "-dBATCH", "-dNOPAUSE", "-dNOOUTERSAVE", "-dNOSAFER", "-sDEVICE=pdfwrite",
                   "-sColorConversionStrategy=RGB", "-dPDFACompatibilityPolicy=1", "-dEmbedAllFonts=true",
                   "-sOutputFile=" + sai, defps, ent]
            try:
                r = subprocess.run(cmd, capture_output=True, timeout=300, creationflags=0x08000000, cwd=tmp)
            except subprocess.TimeoutExpired:
                raise Erro("A conversão de “%s” demorou demais." % nome)
            if r.returncode != 0 or not os.path.exists(sai):
                det = (r.stderr or r.stdout or b"").decode("latin-1", "replace").strip().splitlines()
                raise Erro("Não foi possível converter “%s” para PDF/A%s" % (nome, (": " + det[-1][:160]) if det else "."))
            with open(sai, "rb") as f:
                out = f.read()
            saidas.append((base(nome) + "-pdfa.pdf", out))
            avisos.append("“%s” convertido para PDF/A-%sb (%s)." % (nome, nivel, titulo if titulo else ""))
    avisos.append("Assinaturas digitais, senhas, formulários e anexos não são mantidos no PDF/A. Se precisar assinar, assine depois de converter.")
    return saidas, avisos


# ------------------------------------------------------------------ assinatura digital (certificado A1)
def assinar(arqs, op):
    import datetime
    from pyhanko.sign import signers, fields
    from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
    from pyhanko import stamp
    cert = _img_op(op, "certificado", "Certificado")
    if not cert:
        raise Erro("Escolha o arquivo do certificado digital (.pfx ou .p12).")
    senha = op.get("senha") or ""
    try:
        signer = signers.SimpleSigner.load_pkcs12_data(cert, other_certs=None, passphrase=senha.encode("utf-8"))
    except Exception:
        signer = None
    if signer is None or signer.signing_cert is None:
        raise Erro("Não consegui abrir o certificado: a senha está errada ou o arquivo não é um certificado A1 (.pfx/.p12).")
    cn = signer.signing_cert.subject.native.get("common_name", "") or "Titular do certificado"
    validade = signer.signing_cert["tbs_certificate"]["validity"]["not_after"].native
    agora = datetime.datetime.now(datetime.timezone.utc)
    if validade < agora:
        raise Erro("Este certificado venceu em %s." % validade.strftime("%d/%m/%Y"))
    visivel = op.get("visivel", True)
    pos = op.get("posicao", "baixo-direita")
    saidas, avisos = [], []
    for nome, dados in arqs:
        d = abrir_pdf(dados, nome)
        n = d.page_count
        alvo = n - 1 if op.get("pagina", "ultima") == "ultima" else 0
        R = d[alvo].rect
        ja = sum(1 for _ in d[alvo].widgets() if _.field_type == pymupdf.PDF_WIDGET_TYPE_SIGNATURE) if d.is_pdf else 0
        d.close()
        campo = "Assinatura%d" % (int(time_ns() % 1_000_000))
        w_ = IncrementalPdfFileWriter(io.BytesIO(dados), strict=False)
        meta = signers.PdfSignatureMetadata(field_name=campo, reason=(op.get("motivo") or None), location=(op.get("local") or None),
                                            md_algorithm="sha256", name=cn)
        estilo = None
        if visivel:
            bw, bh, m = 210, 58, 28
            x0 = m if pos.endswith("esquerda") else R.width - m - bw
            y0 = m + ja * (bh + 6) if pos.startswith("baixo") else R.height - m - bh - ja * (bh + 6)
            # pyHanko usa coordenadas do PDF (origem embaixo à esquerda)
            x0, y0 = int(round(x0)), int(round(y0))  # o pyHanko quer números inteiros aqui
            fields.append_signature_field(w_, fields.SigFieldSpec(campo, on_page=alvo, box=(x0, y0, x0 + bw, y0 + bh)))
            # aparência desenhada pelo PyMuPDF (acentos e espaçamento corretos) e usada como carimbo fixo
            linhas = ["Assinado digitalmente por:", cn,
                      "Data: " + datetime.datetime.now().astimezone().strftime("%d/%m/%Y %H:%M:%S %z")]
            if op.get("motivo"):
                linhas.append("Motivo: " + op["motivo"])
            if op.get("local"):
                linhas.append("Local: " + op["local"])
            ap = pymupdf.open()
            apg = ap.new_page(width=bw, height=bh)
            apg.draw_rect(pymupdf.Rect(0.5, 0.5, bw - 0.5, bh - 0.5), color=(0.1, 0.3, 0.6), width=1)
            apg.draw_rect(pymupdf.Rect(0.5, 0.5, 5, bh - 0.5), color=None, fill=(0.1, 0.3, 0.6))
            tam = 7.6 if len(linhas) <= 4 else 6.6
            fonte = pymupdf.Font("helv")
            y = 4 + tam
            for k, t in enumerate(linhas):
                f_nome = "hebo" if k == 1 else "helv"
                while fonte.text_length(t, fontsize=tam) > bw - 16 and len(t) > 4:
                    t = t[:-2] + "…"
                if latin1(t):
                    apg.insert_text((10, y), t, fontsize=tam, fontname=f_nome, color=(0.08, 0.12, 0.2))
                else:
                    apg.insert_text((10, y), t, fontsize=tam, fontname="ASSARIAL", fontfile=fonte_arial(k == 1), color=(0.08, 0.12, 0.2))
                y += tam * 1.32
            ap_path = os.path.join(tempfile.gettempdir(), "carimbo-%s.pdf" % uuid.uuid4().hex)
            ap.save(ap_path)
            ap.close()
            estilo = stamp.StaticStampStyle.from_pdf_file(ap_path, border_width=0)
        else:
            fields.append_signature_field(w_, fields.SigFieldSpec(campo))
        ps = signers.PdfSigner(meta, signer=signer, stamp_style=estilo)
        out = io.BytesIO()
        try:
            ps.sign_pdf(w_, output=out)
        finally:
            if visivel:
                try:
                    os.remove(ap_path)
                except OSError:
                    pass
        saidas.append((base(nome) + "-assinado.pdf", out.getvalue()))
        avisos.append("“%s” assinado com o certificado de %s." % (nome, cn))
    avisos.append("Não altere o arquivo assinado (nem neste editor): qualquer mudança depois da assinatura a invalida.")
    avisos.append("A data da assinatura é a do relógio deste computador (sem carimbo de tempo de servidor).")
    return saidas, avisos


def time_ns():
    import time
    return time.time_ns()


# ------------------------------------------------------------------ registro e execução
# nome -> (função, mínimo de arquivos, máximo, extensões aceitas, precisa do contexto)
PDF = (".pdf",)
IMG = (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff", ".gif")
FERRAMENTAS = {
    "juntar": (juntar, 2, 60, PDF, False), "dividir": (dividir, 1, 10, PDF, False), "remover": (remover, 1, 20, PDF, False),
    "extrair": (extrair, 1, 20, PDF, False), "girar": (girar, 1, 30, PDF, False), "cortar": (cortar, 1, 20, PDF, False),
    "comprimir": (comprimir, 1, 20, PDF, False), "reparar": (reparar, 1, 20, PDF, False), "ocr": (ocr, 1, 10, PDF, True),
    "marca": (marca, 1, 20, PDF, False), "numerar": (numerar, 1, 20, PDF, False), "proteger": (proteger, 1, 20, PDF, False),
    "desbloquear": (desbloquear, 1, 20, PDF, False), "comparar": (comparar, 2, 2, PDF, False),
    "img2pdf": (img2pdf, 1, 100, IMG, False), "office2pdf": (office2pdf, 1, 10, tuple(_OFFICE), False),
    "html2pdf": (html2pdf, 0, 10, (".html", ".htm"), False),
    "pdf2img": (pdf2img, 1, 10, PDF, False), "pdf2word": (pdf2word, 1, 10, PDF, True), "pdf2xlsx": (pdf2xlsx, 1, 10, PDF, False),
    "pdf2pptx": (pdf2pptx, 1, 10, PDF, False), "pdf2txt": (pdf2txt, 1, 20, PDF, False),
    "cabecalho": (cabecalho, 1, 20, PDF, False), "tarja_auto": (tarja_auto, 1, 20, PDF, False),
    "substituir": (substituir, 1, 20, PDF, True), "pdfa": (pdfa, 1, 10, PDF, False), "assinar": (assinar, 1, 10, PDF, False),
    "img_comprimir": (img_comprimir, 1, 100, IMG, False), "img_redimensionar": (img_redimensionar, 1, 100, IMG, False),
    "img_recortar": (img_recortar, 1, 100, IMG, False), "img_converter": (img_converter, 1, 100, IMG, False),
    "img_girar": (img_girar, 1, 100, IMG, False),
}
MIME = {".pdf": "application/pdf", ".zip": "application/zip", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation", ".jpg": "image/jpeg", ".png": "image/png",
        ".webp": "image/webp", ".gif": "image/gif", ".bmp": "image/bmp", ".tiff": "image/tiff",
        ".txt": "text/plain; charset=utf-8"}


def executar(ferramenta, arquivos, opcoes, ctx):
    if ferramenta not in FERRAMENTAS:
        raise Erro("Ferramenta desconhecida.")
    fn, mini, maxi, exts, usa_ctx = FERRAMENTAS[ferramenta]
    if not mini <= len(arquivos) <= maxi:
        raise Erro("Esta ferramenta aceita de %d a %d arquivo(s)." % (mini, maxi) if mini != maxi else "Esta ferramenta precisa de %d arquivo(s)." % mini)
    for nome, dados in arquivos:
        if os.path.splitext(nome.lower())[1] not in exts:
            raise Erro("“%s”: tipo de arquivo não aceito (%s)." % (nome, ", ".join(exts)))
        if not dados:
            raise Erro("“%s” está vazio." % nome)
    saidas, avisos = fn(arquivos, opcoes or {}, ctx) if usa_ctx else fn(arquivos, opcoes or {})
    if not saidas:
        raise Erro("Nada foi gerado.")
    if len(saidas) > 1:
        buf = io.BytesIO()
        usados = set()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for n, b in saidas:
                u, k = n, 1
                while u in usados:
                    k += 1
                    u = "%s (%d)%s" % (os.path.splitext(n)[0], k, os.path.splitext(n)[1])
                usados.add(u)
                z.writestr(u, b)
        saidas = [("%s-resultado.zip" % ferramenta, buf.getvalue())]
    nome, dados = saidas[0]
    return {"nome": nome, "mime": MIME.get(os.path.splitext(nome)[1].lower(), "application/octet-stream"), "bytes": dados, "avisos": avisos}
